import openpyxl
import json
import csv
import os

SPREADSHEET_ID = "1VBMgD1JNHDJxmZV-EwlQC18lwuH8xn3d56RUQS9scts"

def sync_sheets():
    print(f"Starting sync to Google Sheet ID: {SPREADSHEET_ID}")
    
    creds_path = "scripts/google_sheets_credentials.json"
    if not os.path.exists(creds_path):
        print(f"Error: Credentials not found at {creds_path}")
        return False
        
    try:
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
        
        SCOPES = ['https://www.googleapis.com/auth/spreadsheets']
        creds = service_account.Credentials.from_service_account_file(creds_path, scopes=SCOPES)
        service = build('sheets', 'v4', credentials=creds)
        
        # Get existing sheet tabs
        sheet_metadata = service.spreadsheets().get(spreadsheetId=SPREADSHEET_ID).execute()
        existing_sheets = [s['properties']['title'] for s in sheet_metadata.get('sheets', [])]
        print(f"Existing tabs in Google Sheet: {existing_sheets}")
        
        # Read the 3 worksheets from local Excel file
        excel_file = "docs/llm_chatgpt_ads_keyword_analysis.xlsx"
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        
        # Create missing sheets if necessary
        requests = []
        for target_tab in wb.sheetnames:
            if target_tab not in existing_sheets:
                requests.append({
                    'addSheet': {
                        'properties': {
                            'title': target_tab
                        }
                    }
                })
                
        if requests:
            service.spreadsheets().batchUpdate(
                spreadsheetId=SPREADSHEET_ID,
                body={'requests': requests}
            ).execute()
            print(f"Added {len(requests)} new tab(s) to Google Sheet!")
            
        # Update data in each tab
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            
            formatted_data = []
            for r in rows:
                formatted_data.append([str(c) if c is not None else "" for c in r])
                
            # Update range
            service.spreadsheets().values().clear(
                spreadsheetId=SPREADSHEET_ID,
                range=f"'{sheet_name}'!A1:Z50000"
            ).execute()
            
            service.spreadsheets().values().update(
                spreadsheetId=SPREADSHEET_ID,
                range=f"'{sheet_name}'!A1",
                valueInputOption="USER_ENTERED",
                body={'values': formatted_data}
            ).execute()
            
            print(f"✅ Successfully updated sheet '{sheet_name}' with {len(rows)} rows!")
            
        # Remove default 'Sheet1' if present and other tabs exist
        sheet_metadata = service.spreadsheets().get(spreadsheetId=SPREADSHEET_ID).execute()
        updated_sheets = sheet_metadata.get('sheets', [])
        if len(updated_sheets) > 1:
            for s in updated_sheets:
                title = s['properties']['title']
                if title.lower() in ['sheet1', 'sheet 1', 'gid=0'] and len(updated_sheets) > len(wb.sheetnames):
                    sheet_id_to_del = s['properties']['sheetId']
                    try:
                        service.spreadsheets().batchUpdate(
                            spreadsheetId=SPREADSHEET_ID,
                            body={'requests': [{'deleteSheet': {'sheetId': sheet_id_to_del}}]}
                        ).execute()
                        print(f"Cleaned up default sheet '{title}'.")
                    except Exception as ex:
                        pass
                        
        print("\n🎉 GOOGLE SHEET SYNC SUCCESSFUL!")
        print(f"Link: https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit")
        return True
        
    except Exception as e:
        print("Sync Error:", e)
        return False

if __name__ == "__main__":
    sync_sheets()
